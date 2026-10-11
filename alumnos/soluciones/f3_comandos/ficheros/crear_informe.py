import csv
from pptx import Presentation
from pptx.util import Inches, Pt

# ---- Leer el CSV y calcular cifras (nunca de memoria) ----
with open("datos/ventas_tienda.csv", newline="", encoding="utf-8") as f:
    filas = list(csv.DictReader(f))

meses = [fila["mes"] for fila in filas]
hogar_total = sum(int(fila["hogar"]) for fila in filas)
textil_total = sum(int(fila["textil"]) for fila in filas)
bazar_total = sum(int(fila["bazar"]) for fila in filas)
total_general = hogar_total + textil_total + bazar_total

# Verificación: suma de totales mensuales == total general
totales_mensuales = [
    int(fila["hogar"]) + int(fila["textil"]) + int(fila["bazar"]) for fila in filas
]
assert sum(totales_mensuales) == total_general, "Las cifras no cuadran"

# Tendencia: primer mes vs último mes
primer_mes, ultimo_mes = meses[0], meses[-1]
total_primer, total_ultimo = totales_mensuales[0], totales_mensuales[-1]
variacion_pct = (total_ultimo - total_primer) / total_primer * 100

# ---- Crear la presentación (16:9, exactamente 3 diapositivas) ----
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Diapositiva 1: Portada
s1 = prs.slides.add_slide(prs.slide_layouts[0])
s1.shapes.title.text = "Informe semanal de ventas"
s1.placeholders[1].text = ultimo_mes.capitalize()

# Diapositiva 2: Totales por categoría y total general
s2 = prs.slides.add_slide(prs.slide_layouts[5])  # solo título
s2.shapes.title.text = "Totales del periodo"
filas_txt = [
    f"Hogar: {hogar_total:,} €".replace(",", "."),
    f"Textil: {textil_total:,} €".replace(",", "."),
    f"Bazar: {bazar_total:,} €".replace(",", "."),
    f"Total general: {total_general:,} €".replace(",", "."),
]
tb = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
tf = tb.text_frame
for i, linea in enumerate(filas_txt):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = linea
    p.font.size = Pt(24)

# Diapositiva 3: Tendencia
s3 = prs.slides.add_slide(prs.slide_layouts[5])
s3.shapes.title.text = "Tendencia"
texto = (
    f"{primer_mes.capitalize()}: {total_primer:,} €".replace(",", ".")
    + "\n"
    + f"{ultimo_mes.capitalize()}: {total_ultimo:,} €".replace(",", ".")
    + f"\nVariación: {variacion_pct:+.2f} %"
)
tb3 = s3.shapes.add_textbox(Inches(1), Inches(2), Inches(9), Inches(3))
tf3 = tb3.text_frame
for i, linea in enumerate(texto.split("\n")):
    p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
    p.text = linea
    p.font.size = Pt(28)

prs.save("informe_semanal.pptx")

# ---- Verificación final: reabrir y contar diapositivas ----
comprobacion = Presentation("informe_semanal.pptx")
n = len(comprobacion.slides)
print(f"Diapositivas: {n}")
assert n == 3, "Error: la presentación debe tener exactamente 3 diapositivas"
print("OK: 3 diapositivas")
