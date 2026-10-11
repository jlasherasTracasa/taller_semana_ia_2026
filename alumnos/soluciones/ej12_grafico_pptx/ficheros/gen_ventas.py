#!/usr/bin/env python3
"""Genera ventas.pptx leyendo los datos de ventas_tienda.csv.

Los totales se calculan con Python a partir del CSV, nunca a mano.
"""
import csv
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Cm, Pt

BASE = Path(__file__).parent
CSV_PATH = BASE / "ventas_tienda.csv"
PNG_PATH = BASE / "grafico_ventas.png"
PPTX_PATH = BASE / "ventas.pptx"

# ── 1. Leer el CSV ──────────────────────────────────────────────
with open(CSV_PATH, newline="", encoding="utf-8") as f:
    reader = csv.DictReader(f)
    filas = list(reader)

meses = [fila["mes"].capitalize() for fila in filas]
categorias = [c for c in reader.fieldnames if c != "mes"]
datos = {
    cat: [int(fila[cat]) for fila in filas]
    for cat in categorias
}

# ── 2. Totales calculados desde el CSV ──────────────────────────
totales = {cat: sum(vals) for cat, vals in datos.items()}
total_general = sum(totales.values())

print("Totales calculados leyendo", CSV_PATH.name)
for cat, total in totales.items():
    print(f"  {cat:>8}: {total:,} €".replace(",", "."))
print(f"  {'TOTAL':>8}: {total_general:,} €".replace(",", "."))

# ── 3. Gráfico de líneas con matplotlib ─────────────────────────
fig, ax = plt.subplots(figsize=(9, 5), dpi=150)
for cat in categorias:
    ax.plot(meses, datos[cat], marker="o", linewidth=2, label=cat)

ax.set_title("Ventas mensuales por categoría (enero–junio)")
ax.set_xlabel("Mes")
ax.set_ylabel("Ventas (€)")
ax.legend(title="Categoría")
ax.grid(True, linestyle="--", alpha=0.5)
fig.tight_layout()
fig.savefig(PNG_PATH)
plt.close(fig)
print(f"Gráfico guardado en {PNG_PATH.name}")

# ── 4. Presentación de tres diapositivas ────────────────────────
prs = Presentation()
prs.slide_width = Cm(33.87)   # 16:9
prs.slide_height = Cm(19.05)

# Diapositiva 1: portada
portada = prs.slides.add_slide(prs.slide_layouts[0])
portada.shapes.title.text = "Ventas de la tienda"
portada.placeholders[1].text = "Enero – Junio · Resumen por categorías"

# Diapositiva 2: gráfico insertado como imagen
diapo_grafico = prs.slides.add_slide(prs.slide_layouts[5])  # solo título
diapo_grafico.shapes.title.text = "Evolución mensual por categoría"
# Insertar manteniendo la proporción real de la imagen
from PIL import Image
ancho_img = Cm(25.4)
with Image.open(PNG_PATH) as im:
    ratio = im.height / im.width
diapo_grafico.shapes.add_picture(
    str(PNG_PATH),
    left=Cm((33.87 - 25.4) / 2),
    top=Cm(3.2),
    width=ancho_img,
    height=int(ancho_img * ratio),
)

# Diapositiva 3: tabla de totales (calculados del CSV)
diapo_tabla = prs.slides.add_slide(prs.slide_layouts[5])
diapo_tabla.shapes.title.text = "Totales por categoría"

filas_tabla = len(categorias) + 2   # cabecera + categorías + TOTAL GENERAL
tabla_shape = diapo_tabla.shapes.add_table(
    filas_tabla, 2, left=Cm(11.5), top=Cm(4.5), width=Cm(11), height=Cm(1.5) * filas_tabla
)
tabla = tabla_shape.table
tabla.columns[0].width = Cm(7)
tabla.columns[1].width = Cm(4)

def celda(cell, texto, negrita=False):
    cell.text = texto
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(18)
    p.font.bold = negrita

celda(tabla.cell(0, 0), "Categoría", negrita=True)
celda(tabla.cell(0, 1), "Total (€)", negrita=True)
for i, cat in enumerate(categorias, start=1):
    celda(tabla.cell(i, 0), cat.capitalize())
    celda(tabla.cell(i, 1), f"{totales[cat]:,}".replace(",", "."))
fila_total = filas_tabla - 1
celda(tabla.cell(fila_total, 0), "TOTAL GENERAL", negrita=True)
celda(tabla.cell(fila_total, 1), f"{total_general:,}".replace(",", "."), negrita=True)

prs.save(PPTX_PATH)
print(f"Presentación guardada en {PPTX_PATH.name}")
