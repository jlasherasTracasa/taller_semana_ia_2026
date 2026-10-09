"""EJ18: vuelca las facturas PDF (ya extraídas) a facturas.xlsx con openpyxl.

Datos leídos de:
  - facturas/factura_ferreteria.pdf  (Factura 2026/041)
  - facturas/factura_limpieza.pdf    (Factura 2026/019)

Criterio del enunciado: números escritos como número (no fórmula, no texto).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# emisor, nº factura, fecha, base imponible, IVA (21%), total  — todo EUR
FACTURAS = [
    ("Ferretería Etxeberria",  "2026/041", "03/02/2026", 375.60,  78.88,  454.48),
    ("Limpiezas Ribera S.L.",  "2026/019", "28/02/2026", 1002.40, 210.50, 1212.90),
]

wb = Workbook()
ws = wb.active
ws.title = "Facturas"

cabecera = ["Emisor", "Nº factura", "Fecha", "Base imponible (EUR)", "IVA 21% (EUR)", "Total (EUR)"]
ws.append(cabecera)
for c in range(1, len(cabecera) + 1):
    ws.cell(row=1, column=c).font = Font(bold=True)
    ws.cell(row=1, column=c).alignment = Alignment(horizontal="center")

for f in FACTURAS:
    ws.append(list(f))

# Sumas por columna (como NÚMERO, sin fórmulas)
fila_suma = ["SUMA", "", "", sum(f[3] for f in FACTURAS), sum(f[4] for f in FACTURAS), sum(f[5] for f in FACTURAS)]
ws.append(fila_suma)
last = ws.max_row
for c in range(1, len(cabecera) + 1):
    ws.cell(row=last, column=c).font = Font(bold=True)

anchos = [26, 12, 12, 20, 15, 14]
for i, a in enumerate(anchos, start=1):
    ws.column_dimensions[get_column_letter(i)].width = a

out = "facturas.xlsx"
wb.save(out)

# --- Verificación contra los PDF ---
tot_pdf = [454.48, 1212.90]
esperado = round(sum(tot_pdf), 2)
obtenido = round(ws.cell(row=last, column=6).value, 2)
print(f"Totales PDF: {tot_pdf} -> suma esperada: {esperado:.2f}")
print(f"Suma escrita en facturas.xlsx (fila {last}): {obtenido:.2f}")
assert obtenido == esperado == 1667.38, "¡No cuadra!"
print("OK: cuadra con la suma de los totales de los PDF.")
