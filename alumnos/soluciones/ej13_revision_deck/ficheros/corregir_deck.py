#!/usr/bin/env python3
"""Genera deck_ferias_corregido.pptx a partir de deck_ferias.pptx
sin tocar el original. Correcciones puntuales de ortografía,
mayúsculas, espacios y tildes."""
import shutil
from pptx import Presentation

ORIGINAL = "deck_ferias.pptx"
COPIA = "deck_ferias_corregido.pptx"

# Copia del original: el original no se abre para escritura
shutil.copyfile(ORIGINAL, COPIA)

prs = Presentation(COPIA)

# {slide_idx (1-based): {texto_original: texto_corregido}}
CORRECCIONES = {
    1: {
        "El certamen de artesaNIA de Tafalla": "El certamen de artesanía de Tafalla",
    },
    2: {
        "Del 12 al 15 de SEPTIEMBRE": "Del 12 al 15 de septiembre",
        "entrada gratuita": "Entrada gratuita",
    },
    3: {
        "taller de cesteria (sabado 10:00)": "Taller de cestería (sábado 10:00)",
        "CONCIERTO de txistu , domingo": "Concierto de txistu, domingo",
        "Degustación de quesos con D.O. Roncal  ": "Degustación de quesos con D.O. Roncal",
    },
    4: {
        "Organizacion": "Organización",
        "Telefono 948 70 00 00  (extension 3)": "Teléfono 948 70 00 00 (extensión 3)",
    },
}

def iter_runs(prs):
    """Itera sobre todos los runs de texto (shapes y tablas) con su nº de slide."""
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        yield i, run
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                yield i, run

aplicadas = 0
for idx, run in iter_runs(prs):
    reglas = CORRECCIONES.get(idx, {})
    texto = run.text
    if texto in reglas:
        run.text = reglas[texto]
        aplicadas += 1

prs.save(COPIA)
print(f"Correcciones aplicadas: {aplicadas}/{sum(len(v) for v in CORRECCIONES.values())}")
if aplicadas != sum(len(v) for v in CORRECCIONES.values()):
    raise SystemExit("ERROR: no se aplicaron todas las correcciones")
