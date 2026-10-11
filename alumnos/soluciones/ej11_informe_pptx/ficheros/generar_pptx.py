"""Genera presentacion.pptx a partir de informe_cosecha.md (Bodega Larraz, vendimia 2025)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

GRIS = RGBColor(0x59, 0x59, 0x59)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
GRANATE = RGBColor(0x7A, 0x1F, 0x1F)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
DIAPOSITIVA = prs.slide_layouts[6]  # en blanco


def titulo_banda(slide, texto):
    banda = slide.shapes.add_textbox(Inches(0), Inches(0), prs.slide_width, Inches(1.2))
    f = banda.fill
    f.solid()
    f.fore_color.rgb = GRANATE
    p = banda.text_frame.paragraphs[0]
    p.text = texto
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = BLANCO


def caja(slide, izq, arriba, ancho, alto):
    return slide.shapes.add_textbox(Inches(izq), Inches(arriba), Inches(ancho), Inches(alto)).text_frame


# --- 1. Portada -------------------------------------------------------------
s = prs.slides.add_slide(DIAPOSITIVA)
tf = caja(s, 1, 2.5, 11.3, 3)
p = tf.paragraphs[0]
p.text = "Informe de cosecha"
p.font.size = Pt(54)
p.font.bold = True
p2 = tf.add_paragraph()
p2.text = "Bodega Larraz — Vendimia 2025"
p2.font.size = Pt(28)
p2.font.color.rgb = GRIS
p3 = tf.add_paragraph()
p3.text = "Asamblea de la cooperativa · Datos de la vendimia 2025"
p3.font.size = Pt(16)
p3.font.color.rgb = GRIS

# --- 2. Resumen -------------------------------------------------------------
s = prs.slides.add_slide(DIAPOSITIVA)
titulo_banda(s, "Resumen")
tf = caja(s, 1, 1.8, 11.3, 4.8)
p = tf.paragraphs[0]
p.text = ("La vendimia de 2025 cerró con 84.000 kg de uva recogida, un 12% menos que en 2024 "
          "(95.500 kg) por la sequía del verano.")
p.font.size = Pt(24)
p = tf.add_paragraph()
p.text = ("Aun así, la calidad fue excelente: la graduación media subió de 12,1 a 13,4 grados "
          "Baumé.")
p.font.size = Pt(24)

# --- 3. Rendimiento por variedad --------------------------------------------
s = prs.slides.add_slide(DIAPOSITIVA)
titulo_banda(s, "Rendimiento por variedad")
tf = caja(s, 1, 1.8, 11.3, 5)
encabezado = ["Variedad", "Kg", "% del total", "Graduación (Baumé)"]
filas = [
    ["Garnacha", "41.000", "49%", "13,8"],
    ["Tempranillo", "28.500", "34%", "13,1"],
    ["Graciano", "9.500", "11%", "13,0"],
    ["Otras variedades", "5.000", "6%", "12,6"],
]
tbl_shape = s.shapes.add_table(len(filas) + 1, 4, Inches(1), Inches(1.8), Inches(11.3), Inches(4))
tabla = tbl_shape.table
for j, t in enumerate(encabezado):
    c = tabla.cell(0, j)
    c.text = t
    r = c.text_frame.paragraphs[0].runs[0]
    r.font.bold = True
    r.font.size = Pt(20)
for i, fila in enumerate(filas, start=1):
    for j, valor in enumerate(fila):
        c = tabla.cell(i, j)
        c.text = valor
        run = c.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(20)

# --- 4. Evolución ------------------------------------------------------------
s = prs.slides.add_slide(DIAPOSITIVA)
titulo_banda(s, "Evolución (kg totales)")
anios = ["2021", "2022", "2023", "2024", "2025"]
totales = ["71.000", "79.500", "88.000", "95.500", "84.000"]
ancho = 2.1
for i, (anio, kg) in enumerate(zip(anios, totales)):
    x = 0.8 + i * (ancho + 0.4)
    anio_box = caja(s, x, 2.2, ancho, 0.6)
    p = anio_box.paragraphs[0]
    p.text = anio
    p.font.size = Pt(24)
    p.font.bold = True
    kg_box = caja(s, x, 2.9, ancho, 0.8)
    p = kg_box.paragraphs[0]
    p.text = kg + " kg"
    p.font.size = Pt(22)
    p.font.color.rgb = GRIS

nota = caja(s, 1, 5.6, 11.3, 1)
p = nota.paragraphs[0]
p.text = "Máximo en 2024 (95.500 kg); retroceso del 12% en 2025 por la sequía."
p.font.size = Pt(18)
p.font.italic = True

# --- 5. Conclusiones ---------------------------------------------------------
s = prs.slides.add_slide(DIAPOSITIVA)
titulo_banda(s, "Conclusiones")
tf = caja(s, 1, 1.8, 11.3, 4.5)
p = tf.paragraphs[0]
p.text = ("Menos cantidad pero más grado: se prevé un vino estructurado con mayor potencial de "
          "crianza.")
p.font.size = Pt(24)
p = tf.add_paragraph()
p.text = "Se recomienda ampliar la parcela de Garnacha en 2026."
p.font.size = Pt(24)

prs.save("presentacion.pptx")
print(f"OK: presentacion.pptx con {len(prs.slides._sldIdLst)} diapositivas 16:9")
