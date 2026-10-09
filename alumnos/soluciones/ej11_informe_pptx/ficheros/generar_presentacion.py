# -*- coding: utf-8 -*-
"""Genera presentacion.pptx a partir de informe_cosecha.md."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ACCENT = RGBColor(0x6B, 0x1F, 0x2F)   # burdeos sobrio
DARK = RGBColor(0x33, 0x33, 0x33)

prs = Presentation()
prs.slide_width = Inches(13.333)      # 16:9
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def add_text(slide, text, left, top, width, height, size=18, bold=False,
             color=DARK):
    box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_title(slide, text):
    add_text(slide, text, 0.8, 0.6, 11.7, 0.9, size=30, bold=True, color=ACCENT)


# 1 · Portada
s = prs.slides.add_slide(blank)
add_text(s, "Bodega Larraz", 1.0, 2.3, 11.3, 1.0, size=44, bold=True, color=ACCENT)
add_text(s, "Informe de cosecha — Vendimia 2025", 1.0, 3.4, 11.3, 0.8, size=24)

# 2 · Resumen
s = prs.slides.add_slide(blank)
add_title(s, "Resumen de la vendimia 2025")
add_text(s,
         "• La vendimia de 2025 cerró con 84.000 kg de uva recogida.\n"
         "• Un 12 % menos que en 2024 (95.500 kg) por la sequía del verano.\n"
         "• La calidad fue excelente: la graduación media subió de\n"
         "   12,1 a 13,4 grados Baumé.",
         0.8, 1.9, 11.7, 3.5, size=20)

# 3 · Rendimiento por variedad
s = prs.slides.add_slide(blank)
add_title(s, "Rendimiento por variedad")
rows = [
    ("Variedad", "Kg recogidos", "% del total", "Graduación (Baumé)"),
    ("Garnacha", "41.000 kg", "49 %", "13,8"),
    ("Tempranillo", "28.500 kg", "34 %", "13,1"),
    ("Graciano", "9.500 kg", "11 %", "13,0"),
    ("Otras variedades", "5.000 kg", "6 %", "12,6"),
]
tbl = s.shapes.add_table(5, 4, Inches(0.8), Inches(1.9),
                         Inches(11.7), Inches(3.8)).table
tbl.columns[0].width = Inches(3.5)
tbl.columns[1].width = Inches(2.6)
tbl.columns[2].width = Inches(2.2)
tbl.columns[3].width = Inches(3.4)
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.text = val
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(16)
        run.font.bold = (i == 0)
        if i == 0:
            run.font.color.rgb = ACCENT

# 4 · Evolución
s = prs.slides.add_slide(blank)
add_title(s, "Evolución (kg totales)")
years = [("2021", "71.000"), ("2022", "79.500"), ("2023", "88.000"),
         ("2024", "95.500"), ("2025", "84.000")]
for i, (year, kg) in enumerate(years):
    x = 0.9 + i * 2.4
    add_text(s, year, x, 2.2, 2.0, 0.6, size=18, bold=True, color=ACCENT)
    add_text(s, kg + " kg", x, 3.0, 2.0, 0.6, size=22)

# 5 · Conclusiones
s = prs.slides.add_slide(blank)
add_title(s, "Conclusiones")
add_text(s,
         "• Menos cantidad pero más grado: se prevé un vino estructurado\n"
         "   con mayor potencial de crianza.\n"
         "• Se recomienda ampliar la parcela de Garnacha en 2026.",
         0.8, 1.9, 11.7, 3.0, size=20)

prs.save("presentacion.pptx")

# Verificación: releer el archivo generado
check = Presentation("presentacion.pptx")
print("Diapositivas:", len(check.slides))
